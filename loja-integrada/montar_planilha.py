# -*- coding: utf-8 -*-
"""Preenche o modelo oficial da Loja Integrada com os 20 produtos.

Usa o python do sistema (tem openpyxl); o venv do projeto não tem.

    python -I loja-integrada/montar_planilha.py

Nome, preço e descrição vêm de `tray-produtos.json`, raspado da loja Tray em
13/10/2026. Peso e dimensões são ESTIMADOS, menos o 1 kg; o LEIA-ME diz quais.
"""
import os, sys, json, re, shutil
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
from catalogo import PRODUTOS, FORMATO

MODELO = os.path.join(AQUI, "planilha-modelo.xlsx")
SAIDA = os.path.join(AQUI, "entrega", "planilha-produtos-cafe-canastra.xlsx")
NCM_CAFE = "09012100"  # café torrado, não descafeinado — confirmar com a contabilidade


def limpa_nome(n):
    return re.sub(r"\s*-\s*Café Canastra\s*$", "", n).strip()


def html(desc):
    paragrafos = [p.strip() for p in desc.split("\n\n") if p.strip()]
    return "".join(f"<p>{p}</p>" for p in paragrafos)


def seo_descricao(desc):
    plano = " ".join(desc.split())
    return plano[:157].rsplit(" ", 1)[0] + "..." if len(plano) > 160 else plano


def main():
    tray = {p["sku"]: p for p in json.load(open(os.path.join(AQUI, "tray-produtos.json"), encoding="utf-8"))}
    fotos = json.load(open(os.path.join(AQUI, "mapa-fotos.json"), encoding="utf-8"))

    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    shutil.copy(MODELO, SAIDA)
    wb = openpyxl.load_workbook(SAIDA)
    ws = wb.active
    col = {c.value: c.column for c in ws[1]}
    ws.delete_rows(2, ws.max_row)  # fora as linhas de exemplo

    for i, (sku, tsku, _pg, _pr, cat, fmt, ean) in enumerate(PRODUTOS, start=2):
        p = tray[tsku]
        peso, alt, larg, comp = FORMATO[fmt]
        nome = limpa_nome(p["nome"])
        marca = "Néctar de Minas" if sku.startswith("NEC") else "Café Canastra"
        valores = {
            "tipo": "sem-variacao", "sku": sku, "ativo": "S", "usado": "N", "destaque": "N",
            "ncm": NCM_CAFE, "gtin": ean, "nome": nome,
            "seo-tag-title": nome[:70],
            "seo-tag-description": seo_descricao(p["desc_full"]),
            "descricao-completa": html(p["desc_full"]),
            "estoque-gerenciado": "N",
            "estoque-situacao-em-estoque": "imediata",
            "estoque-situacao-sem-estoque": "indisponivel",
            "preco-sob-consulta": "N",
            "preco-cheio": float(p["preco"]),
            "marca": marca,
            "peso-em-kg": peso, "altura-em-cm": alt, "largura-em-cm": larg, "comprimento-em-cm": comp,
            "categoria-nome-nivel-1": cat[0], "categoria-nome-nivel-2": cat[1],
        }
        for chave, valor in valores.items():
            ws.cell(row=i, column=col[chave], value=valor)
        # imagem-1..5 ficam VAZIAS: a entrega é de arquivo, não de URL (decisão do cliente).
        assert len(fotos[sku]) >= 1, f"{sku} sem foto"

    wb.save(SAIDA)
    print(f"{ws.max_row - 1} linhas · {SAIDA}")
    # releitura de conferência
    wb2 = openpyxl.load_workbook(SAIDA)
    ws2 = wb2.active
    for linha in ws2.iter_rows(min_row=2, values_only=True):
        d = dict(zip([c.value for c in ws2[1]], linha))
        print(f"{d['sku']:<16} R$ {d['preco-cheio']:>7.2f}  {d['peso-em-kg']:>5} kg  "
              f"gtin {str(d['gtin'] or '—'):<14} {d['nome'][:52]}")


if __name__ == "__main__":
    sys.exit(main())
